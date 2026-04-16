"""Extract text and vision payloads from user-uploaded files for the AI assistant.

Supports PDF, DOCX, XLSX/CSV as text, and common image formats as base64 for vision.
"""

from __future__ import annotations

import base64
import io
import logging
from dataclasses import dataclass, field
from typing import Any

from fastapi import UploadFile

logger = logging.getLogger(__name__)

# Total extracted text budget (chars) per request to stay within token limits.
MAX_COMBINED_TEXT_CHARS = 80_000
MAX_SINGLE_FILE_BYTES = 12 * 1024 * 1024
MAX_FILES = 12
MAX_IMAGES = 8

IMAGE_MIME_SET = frozenset(
    ("image/jpeg", "image/jpg", "image/png", "image/gif", "image/webp"),
)


@dataclass
class AttachmentBuildResult:
    """OpenAI multimodal content parts plus a short line for chat history."""

    content_parts: list[dict[str, Any]]
    history_line: str
    warnings: list[str] = field(default_factory=list)
    file_summaries: list[dict[str, Any]] = field(default_factory=list)
    vision_image_count: int = 0
    total_document_text_chars: int = 0
    #: Original ``.xlsx`` bytes keyed by upload filename (for server-side import tool injection).
    xlsx_raw_by_filename: dict[str, bytes] = field(default_factory=dict)


def format_context_acknowledgement_line(summaries: list[dict[str, Any]]) -> str:
    """Short Russian line listing which files reached the model context (for chat history)."""
    if not summaries:
        return ""
    lines: list[str] = []
    for s in summaries:
        name = s.get("name") or "file"
        if s.get("included_in_context"):
            if (s.get("kind") or "") == "image":
                lines.append(f"• {name} — изображение передано в vision")
            elif s.get("text_chars") is not None:
                lines.append(f"• {name} — текст {int(s['text_chars'])} симв. в контексте")
            else:
                lines.append(f"• {name} — в контексте")
        else:
            note = (s.get("note") or "не попал в запрос").strip()
            lines.append(f"• {name} — ⚠ {note}")
    return "Контекст модели:\n" + "\n".join(lines)


def _truncate(text: str, limit: int) -> tuple[str, bool]:
    if len(text) <= limit:
        return text, False
    return text[: limit - 3] + "...", True


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for i, page in enumerate(reader.pages[:80]):
        try:
            t = page.extract_text() or ""
        except Exception:
            t = ""
        if t.strip():
            parts.append(f"--- Page {i + 1} ---\n{t}")
    return "\n\n".join(parts).strip()


def _extract_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append("\t".join(cells))
    return "\n".join(lines).strip()


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: list[str] = []
    for sheet in wb.worksheets:
        out.append(f"## {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            out.append("\t".join("" if c is None else str(c) for c in row))
    return "\n".join(out).strip()


def _extract_csv(data: bytes) -> str:
    return data.decode("utf-8-sig", errors="replace").strip()


def _guess_kind(filename: str, content_type: str | None) -> str:
    ct = (content_type or "").split(";")[0].strip().lower()
    fn = (filename or "").lower()
    if ct.startswith("image/") or fn.endswith((".jpg", ".jpeg", ".png", ".gif", ".webp")):
        return "image"
    if ct == "application/pdf" or fn.endswith(".pdf"):
        return "pdf"
    if (
        ct == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or fn.endswith(".docx")
    ):
        return "docx"
    if (
        ct == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        or fn.endswith(".xlsx")
    ):
        return "xlsx"
    if ct.startswith("text/csv") or fn.endswith(".csv"):
        return "csv"
    if fn.endswith(".doc") and not fn.endswith(".docx"):
        return "doc_legacy"
    raise ValueError(f"Unsupported file type: {filename} ({content_type})")


async def build_upload_parts(
    instruction: str,
    files: list[UploadFile],
) -> AttachmentBuildResult:
    """Build OpenAI multimodal user content and a human-readable history line."""

    warnings: list[str] = []
    if len(files) > MAX_FILES:
        raise ValueError(f"At most {MAX_FILES} files per message")

    text_parts: list[str] = []
    image_parts: list[dict[str, Any]] = []
    labels: list[str] = []
    file_summaries: list[dict[str, Any]] = []
    xlsx_raw_by_filename: dict[str, bytes] = {}

    for uf in files:
        if not uf.filename:
            continue
        raw = await uf.read()
        if len(raw) > MAX_SINGLE_FILE_BYTES:
            raise ValueError(
                f"File «{uf.filename}» exceeds {MAX_SINGLE_FILE_BYTES // (1024 * 1024)} MB",
            )

        kind = _guess_kind(uf.filename, uf.content_type)
        labels.append(uf.filename)
        if kind == "xlsx":
            xlsx_raw_by_filename[uf.filename] = raw

        if kind == "image":
            mime = (uf.content_type or "image/jpeg").split(";")[0].strip().lower()
            if mime not in IMAGE_MIME_SET and not mime.startswith("image/"):
                mime = "image/jpeg"
            image_parts.append(
                {
                    "type": "image_url",
                    "image_url": {
                        "url": f"data:{mime};base64,{base64.standard_b64encode(raw).decode('ascii')}",
                    },
                }
            )
            file_summaries.append(
                {
                    "name": uf.filename,
                    "kind": "image",
                    "included_in_context": True,
                    "note": None,
                }
            )
            continue

        if kind == "doc_legacy":
            warnings.append(
                f"Файл «{uf.filename}» в формате .doc не поддерживается. Сохраните как .docx или PDF.",
            )
            file_summaries.append(
                {
                    "name": uf.filename,
                    "kind": "doc_legacy",
                    "included_in_context": False,
                    "note": "формат .doc не поддерживается",
                }
            )
            continue

        try:
            if kind == "pdf":
                extracted = _extract_pdf(raw)
            elif kind == "docx":
                extracted = _extract_docx(raw)
            elif kind == "xlsx":
                extracted = _extract_xlsx(raw)
            elif kind == "csv":
                extracted = _extract_csv(raw)
            else:
                extracted = ""
        except Exception as exc:
            logger.warning("attachment extract failed", extra={"file": uf.filename, "error": str(exc)})
            warnings.append(f"Не удалось прочитать «{uf.filename}»: {exc}")
            file_summaries.append(
                {
                    "name": uf.filename,
                    "kind": kind,
                    "included_in_context": False,
                    "note": f"ошибка чтения: {exc!s}"[:500],
                }
            )
            continue

        if not extracted.strip():
            warnings.append(
                f"Файл «{uf.filename}»: текст не извлечён (возможно, скан). "
                "Загрузите страницы как изображения (JPEG/PNG).",
            )
            file_summaries.append(
                {
                    "name": uf.filename,
                    "kind": kind,
                    "included_in_context": False,
                    "note": "текст не извлечён (скан?)",
                }
            )
            continue

        text_parts.append(f"### Файл: {uf.filename}\n\n{extracted}")
        file_summaries.append(
            {
                "name": uf.filename,
                "kind": kind,
                "included_in_context": True,
                "text_chars": len(extracted),
                "note": None,
            }
        )

    if len(image_parts) > MAX_IMAGES:
        raise ValueError(f"At most {MAX_IMAGES} images per message")

    if labels and not text_parts and not image_parts:
        raise ValueError(
            "Не удалось извлечь данные из вложений. Для сканов используйте изображения (JPEG/PNG) "
            "или укажите текстовый PDF/DOCX/XLSX.",
        )

    combined_text = "\n\n".join(text_parts)
    combined_text, truncated = _truncate(combined_text, MAX_COMBINED_TEXT_CHARS)
    if truncated:
        warnings.append(
            "Текст вложений обрезан по объёму; при необходимости разбейте на несколько сообщений.",
        )

    instr = (instruction or "").strip()
    if not instr and not text_parts and not image_parts:
        raise ValueError("Добавьте сообщение или файлы")

    default_instr = (
        "Проанализируй вложения (ТЗ, ведомость работ, таблица или сканы). "
        "Извлеки отдельные работы/пункты и создай соответствующие задачи в CRM. "
        "Для списка строк используй bulk_create_tasks; для отдельных пунктов — create_task. "
        "Сопоставь колонки: название работы → title; срок/дата → due_date; ответственный → assignee_query или assigned_to."
    )
    if not instr:
        instr = default_instr

    block = instr
    if combined_text:
        block = f"{instr}\n\n--- Содержимое документов ---\n\n{combined_text}"

    content_parts: list[dict[str, Any]] = [{"type": "text", "text": block}]
    content_parts.extend(image_parts)

    ack = format_context_acknowledgement_line(file_summaries)
    hist = instr[:2000]
    if labels:
        hist = f"{hist}\n\n[Вложения: {', '.join(labels)}]" if hist else f"[Вложения: {', '.join(labels)}]"
    if ack:
        hist = f"{ack}\n\n{hist}" if hist.strip() else ack
    if warnings:
        hist += "\n\n" + " ".join(warnings[:3])

    return AttachmentBuildResult(
        content_parts=content_parts,
        history_line=hist.strip() or "[Файлы без текста]",
        warnings=warnings,
        file_summaries=file_summaries,
        vision_image_count=len(image_parts),
        total_document_text_chars=len(combined_text),
        xlsx_raw_by_filename=xlsx_raw_by_filename,
    )
