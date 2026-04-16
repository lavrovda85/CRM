"""Classify Excel sheets and import clients and/or warehouse rows."""

from __future__ import annotations

import uuid
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.schemas.excel_import import ExcelUnifiedImportResponse, SheetImportSummary
from app.services.excel_client_import import (
    find_client_header_row,
    import_clients_from_sheet,
    score_client_header_row,
)
from app.services.warehouse_xlsx_import import import_warehouse_rows


def _header_line(row: tuple[object, ...] | None) -> str:
    if not row:
        return ""
    parts: list[str] = []
    for c in row:
        if c is None:
            continue
        parts.append(str(c).replace("\xa0", " ").strip().lower())
    return " ".join(parts)


def _warehouse_sheet_score(rows: list[tuple[object, ...]]) -> int:
    """Heuristic: stock sheet has SKU + quantity style headers."""
    if not rows:
        return 0
    best = 0
    for i in range(min(5, len(rows))):
        line = _header_line(rows[i])
        sc = 0
        if "артикул" in line or " sku" in line or line.strip().startswith("sku"):
            sc += 3
        if "количество" in line or "кол-во" in line:
            sc += 2
        if "название" in line or "наименование" in line:
            sc += 1
        if "ед" in line and "изм" in line:
            sc += 1
        if "мин" in line and "остат" in line:
            sc += 1
        best = max(best, sc)
    return best


def _client_sheet_score(rows: list[tuple[object, ...]]) -> int:
    if not rows:
        return 0
    best = 0
    for i in range(min(15, len(rows))):
        row = rows[i]
        if not row:
            continue
        best = max(best, score_client_header_row(row))
    return best


def classify_sheet(rows: list[tuple[object, ...]]) -> str:
    """Return ``clients``, ``warehouse``, or ``unknown``."""
    wh = _warehouse_sheet_score(rows)
    cl = _client_sheet_score(rows)
    if wh >= 4 and cl < 4:
        return "warehouse"
    if cl >= 4 and wh < 4:
        return "clients"
    if cl >= wh and cl >= 3:
        return "clients"
    if wh >= 3:
        return "warehouse"
    return "unknown"


async def import_excel_workbook(
    db: AsyncSession,
    *,
    company_id: uuid.UUID,
    raw_bytes: bytes,
    use_ai_hint: bool = False,
) -> ExcelUnifiedImportResponse:
    """Import all worksheets: each sheet classified as client list or warehouse table.

    Args:
        db: DB session.
        company_id: Active tenant.
        raw_bytes: XLSX file bytes.
        use_ai_hint: Reserved for future AI-assisted column mapping (no-op for now).

    Returns:
        Aggregated counts and per-sheet notes.
    """
    _ = use_ai_hint  # future: call OpenAI when classification is ambiguous

    try:
        wb = load_workbook(filename=BytesIO(raw_bytes), data_only=True)
    except Exception as exc:
        return ExcelUnifiedImportResponse(
            errors=[f"Failed to read XLSX: {exc}"],
        )

    total_cc = 0
    total_cs = 0
    total_wc = 0
    total_wu = 0
    total_ws = 0
    sheets_out: list[SheetImportSummary] = []
    all_errors: list[str] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            sheets_out.append(
                SheetImportSummary(sheet_name=sheet_name, kind="skipped", message="empty")
            )
            continue

        kind = classify_sheet(rows)
        if kind == "warehouse":
            res = await import_warehouse_rows(db, company_id=company_id, rows=rows)
            total_wc += res.created_count
            total_wu += res.updated_count
            total_ws += res.skipped_count
            all_errors.extend(res.errors or [])
            sheets_out.append(
                SheetImportSummary(
                    sheet_name=sheet_name,
                    kind="warehouse",
                    rows_processed=res.created_count + res.updated_count + res.skipped_count,
                    message=f"+{res.created_count} new, {res.updated_count} updated",
                )
            )
        elif kind == "clients":
            cc, cs, errs = await import_clients_from_sheet(
                db,
                company_id=company_id,
                rows=rows,
                sheet_name=sheet_name,
            )
            total_cc += cc
            total_cs += cs
            all_errors.extend(errs)
            sheets_out.append(
                SheetImportSummary(
                    sheet_name=sheet_name,
                    kind="clients",
                    rows_processed=cc + cs,
                    message=f"{cc} created, {cs} skipped",
                )
            )
        else:
            # Try client import anyway if any row looks like data
            idx = find_client_header_row(rows)
            if idx is not None:
                cc, cs, errs = await import_clients_from_sheet(
                    db,
                    company_id=company_id,
                    rows=rows,
                    sheet_name=sheet_name,
                )
                total_cc += cc
                total_cs += cs
                all_errors.extend(errs)
                sheets_out.append(
                    SheetImportSummary(
                        sheet_name=sheet_name,
                        kind="clients",
                        rows_processed=cc + cs,
                        message=f"low-confidence mapping: {cc} created, {cs} skipped",
                    )
                )
            else:
                sheets_out.append(
                    SheetImportSummary(
                        sheet_name=sheet_name,
                        kind="skipped",
                        message="could not detect client or warehouse headers",
                    )
                )

    wb.close()
    return ExcelUnifiedImportResponse(
        clients_created=total_cc,
        clients_skipped=total_cs,
        warehouse_created=total_wc,
        warehouse_updated=total_wu,
        warehouse_skipped=total_ws,
        sheets=sheets_out,
        ai_mapping_used=False,
        errors=all_errors[:50],
    )
