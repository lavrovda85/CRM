"""Normalize client workbook columns with pandas-first heuristics.

This script reads each sheet into a pandas DataFrame for primary row analysis,
merges row cells into one semantic text block, extracts normalized company,
contact, address, phones, and email values, and writes the results back into
new auto-columns in the workbook.
"""

from __future__ import annotations

import re
import shutil
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

LEGAL_MARKERS = (
    "ооо",
    "оао",
    "зао",
    "пао",
    "ао",
    "ип",
    "гбу",
    "мбу",
    "фгбоу",
    "фгау",
    "муп",
    "кгау",
    "филиал",
    "ltd",
    "llc",
    "inc",
)

ROLE_WORDS = {
    "ведущий",
    "главный",
    "генеральный",
    "менеджер",
    "юрист",
    "инженер",
    "директор",
    "специалист",
    "логист",
    "начальник",
    "администратор",
    "руководитель",
    "бухгалтер",
    "секретарь",
}

GENERIC_NON_COMPANY_WORDS = {
    "физ",
    "физлицо",
    "физ.лицо",
    "физ лицо",
    "юр",
    "юрлицо",
    "юр.лицо",
    "юр лицо",
    "город",
    "адрес",
    "красноярск",
    "коттедж",
}

BUSINESS_HINT_WORDS = {
    "салон",
    "студия",
    "магазин",
    "кафе",
    "ресторан",
    "гостиница",
    "отель",
    "клиника",
    "аптека",
    "банк",
    "сервис",
    "центр",
    "комфорт",
    "снабжение",
    "торг",
    "мир",
    "красоты",
}

ADDRESS_WORD_BLACKLIST = {
    "лето",
    "зима",
    "весна",
    "осень",
    "работал",
    "обслуживание",
    "контрагент",
    "кондиционер",
    "коттедж",
    "chigo",
    "electroluxe",
}

COMMON_FIRST_NAMES = {
    "александр",
    "алексей",
    "анатолий",
    "андрей",
    "анна",
    "анфиса",
    "арина",
    "вадим",
    "валентина",
    "валерий",
    "василий",
    "виктор",
    "виталий",
    "владимир",
    "вячеслав",
    "галина",
    "дарья",
    "денис",
    "дмитрий",
    "евгений",
    "екатерина",
    "елена",
    "игорь",
    "илья",
    "ирина",
    "ксения",
    "лариса",
    "любовь",
    "максим",
    "марина",
    "мария",
    "михаил",
    "надежда",
    "наталья",
    "николай",
    "олег",
    "ольга",
    "павел",
    "петр",
    "роман",
    "светлана",
    "сергей",
    "татьяна",
    "юлия",
    "яна",
}

RUSSIAN_NAME_ENDINGS = (
    "ов",
    "ова",
    "ев",
    "ева",
    "ин",
    "ина",
    "ын",
    "ына",
    "ский",
    "ская",
    "цкий",
    "цкая",
    "ко",
    "ук",
    "юк",
    "ич",
    "вич",
    "вна",
    "ична",
    "ович",
    "евич",
    "ьевич",
    "оглы",
    "ян",
    "янц",
    "ия",
    "ий",
    "ей",
    "ай",
    "ан",
    "ена",
    "ита",
    "иля",
    "ья",
)

ADDR_TOKEN_RE = re.compile(
    r"\b(ул\.?|улица|пр\.?|просп\.?|проспект|пер\.?|переулок|ш\.?|шоссе|б-р|бульвар|пл\.?|площадь|наб\.?|набережная|проезд|пр-д|мкр\.?|м-н|дом|д\.?|оф\.?|офис|строение|стр\.?|корп\.?|корпус|кв\.?|квартира|пом\.?|помещение|район|р-н)\b",
    flags=re.IGNORECASE,
)
EMAIL_RE = re.compile(r"(?i)\b[a-z0-9._%+\-]+@[a-z0-9.\-]+\.[a-z]{2,}\b")
PHONE_BLOCK_RE = re.compile(r"(?<!\w)(?:\+?\d[\d\s()\-]{4,}\d)(?!\w)")
PHONE_DIGIT_TOKEN_RE = re.compile(r"(?<![\d/\-])(?:\+?7|8)?\d{10}(?!\d)|(?<![\d/\-])\d{6}(?!\d)")
PHONE_FLEX_RE = re.compile(r"(?<![\d/\-])(?:\+?7|8)[\d()\-\s]{9,20}\d(?!\d)")
DATEISH_RE = re.compile(r"^(?:\d{1,2}[.,\-/]\d{1,2}(?:[.,\-/]\d{2,4})?|\d{4}[.,\-/]\d{1,2}[.,\-/]\d{1,2})$")
PERSON_RE = re.compile(r"\b[А-ЯЁ][а-яё]+(?:\s+[А-ЯЁ][а-яё]+){1,2}\b")
INITIALS_RE = re.compile(r"\b[А-ЯЁ][а-яё]+\s+[А-ЯЁ]\.\s*[А-ЯЁ]\.(?:\s+[А-ЯЁ][а-яё]+)?\b")
SINGLE_NAME_RE = re.compile(r"\b[А-ЯЁ][а-яё]{2,}\b")
CITY_RE = re.compile(r"\b(?:г\.?\s*)?[А-ЯЁ][а-яё]+(?:-[А-ЯЁ][а-яё]+)?\b")
STREET_WITH_NUMBER_RE = re.compile(
    r"\b(?:[А-ЯЁ][а-яё\-]+\s+){1,3}\d+[а-яёa-z]?\b(?:\s*(?:оф\.?|кв\.?|пом\.?|стр\.?|корп\.?|д\.?|дом)\s*\d+[а-яёa-z]?)?",
    flags=re.IGNORECASE,
)
COMPACT_ADDRESS_RE = re.compile(
    r"(?:\bг\.?\s*[А-ЯЁ][а-яё-]+\s*,\s*)?"
    r"\b[А-ЯЁ][а-яё-]+(?:\s+[А-ЯЁ][а-яё-]+){0,2}\s*,\s*"
    r"\d+[а-яёa-z]?(?:/\d+[а-яёa-z]?)?(?:-\d+)?"
    r"(?:\s*(?:кв\.?|оф\.?|пом\.?)\s*\d+[а-яёa-z]?)?",
    flags=re.IGNORECASE,
)
SPACE_RE = re.compile(r"\s+")
QUOTE_EDGE_RE = re.compile(r'^["«»\']+|["«»\']+$')

NEW_COLUMNS = [
    "Компания (auto)",
    "Контактное лицо (auto)",
    "Адрес (auto)",
    "Телефоны (auto)",
    "Email (auto)",
]


@dataclass(slots=True)
class RowNormalizationResult:
    """Normalized row values."""

    company: str
    contact: str
    address: str
    phones: str
    emails: str


def clean_text(value: object) -> str:
    text = "" if value is None else str(value)
    text = text.replace("_x000D_", " ").replace("\xa0", " ")
    return SPACE_RE.sub(" ", text.replace("\n", " ")).strip(" ,;\t")


def dedupe(items: Iterable[str]) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for item in items:
        x = clean_text(item)
        if not x:
            continue
        key = x.lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(x)
    return out


def normalize_company_text(text: str) -> str:
    src = clean_text(text)
    replacements = {
        "ооо": "ООО",
        "оао": "ОАО",
        "зао": "ЗАО",
        "пао": "ПАО",
        "ао": "АО",
        "ип": "ИП",
    }
    for low, up in replacements.items():
        src = re.sub(rf"(?i)\b{re.escape(low)}\b", up, src)
    src = QUOTE_EDGE_RE.sub("", src)
    src = src.replace("“", '"').replace("”", '"')
    src = re.sub(r'\s+"', ' "', src)
    return clean_text(src)


def has_legal_marker(text: str) -> bool:
    low = clean_text(text).lower()
    return any(marker in low for marker in LEGAL_MARKERS)


def is_russian_name_token(token: str) -> bool:
    t = clean_text(token)
    if not t or not re.fullmatch(r"[А-ЯЁ][а-яё]+", t):
        return False
    low = t.lower()
    if low in GENERIC_NON_COMPANY_WORDS or low in ROLE_WORDS:
        return False
    return low in COMMON_FIRST_NAMES or low.endswith(RUSSIAN_NAME_ENDINGS)


def has_street_number_candidate(text: str) -> bool:
    phrase_match = STREET_WITH_NUMBER_RE.search(clean_text(text))
    if not phrase_match:
        return False
    phrase = clean_text(phrase_match.group(0))
    head = re.sub(r"\d.*$", "", phrase).strip(" ,")
    words = [word.lower() for word in re.findall(r"[А-ЯЁа-яёA-Za-z-]+", head)]
    if not words:
        return False
    if any(word in ADDRESS_WORD_BLACKLIST or word in BUSINESS_HINT_WORDS for word in words):
        return False
    if len(words) == 1 and words[0] in COMMON_FIRST_NAMES:
        return False
    if len(words) == 1:
        single = words[0]
        return len(single) >= 5 and (
            single.endswith(("ина", "ова", "ева", "ского", "ский", "ая", "яя", "иная"))
            or single not in COMMON_FIRST_NAMES
        )
    adjective_like = any(
        word.endswith(("ой", "ая", "яя", "ий", "ый", "ого", "его", "ской", "ский"))
        for word in words
    )
    return adjective_like or len(words) >= 2


def looks_like_address(text: str) -> bool:
    t = clean_text(text)
    if not t:
        return False
    if ADDR_TOKEN_RE.search(t) and re.search(r"\d", t):
        return True
    if COMPACT_ADDRESS_RE.search(t):
        return True
    return has_street_number_candidate(t)


def looks_like_person(text: str) -> bool:
    t = clean_text(text)
    if not t or has_legal_marker(t) or looks_like_address(t):
        return False
    if INITIALS_RE.fullmatch(t) or PERSON_RE.fullmatch(t):
        return True
    parts = t.split()
    return 1 <= len(parts) <= 3 and all(is_russian_name_token(part) for part in parts)


def format_phone(raw: str) -> str:
    digits = re.sub(r"\D", "", raw)
    if len(digits) == 11 and digits[0] in ("7", "8"):
        d = "7" + digits[1:]
        return f"+{d[0]} ({d[1:4]}) {d[4:7]}-{d[7:9]}-{d[9:11]}"
    if len(digits) == 10:
        return f"+7 ({digits[0:3]}) {digits[3:6]}-{digits[6:8]}-{digits[8:10]}"
    if len(digits) == 6:
        return f"{digits[:3]}-{digits[3:]}"
    return clean_text(raw)


def extract_emails(text: str) -> list[str]:
    return dedupe(match.group(0).lower() for match in EMAIL_RE.finditer(text))


def extract_phones(text: str) -> list[str]:
    out: list[str] = []
    seen_raw: list[str] = [match.group(0) for match in PHONE_BLOCK_RE.finditer(text)]
    seen_raw.extend(match.group(0) for match in PHONE_DIGIT_TOKEN_RE.finditer(text))
    seen_raw.extend(match.group(0) for match in PHONE_FLEX_RE.finditer(text))
    for raw_match in seen_raw:
        raw = clean_text(raw_match)
        compact = raw.replace(" ", "")
        digits = re.sub(r"\D", "", raw)
        if DATEISH_RE.fullmatch(compact):
            continue
        if re.search(r"\d{4}[\-/]\d{2}[\-/]\d{2}", compact):
            continue
        if len(digits) == 11 and digits[0] in ("7", "8"):
            out.append(format_phone(raw))
        elif len(digits) == 10 and digits[0] in ("3", "4", "8", "9"):
            out.append(format_phone(raw))
        elif len(digits) == 6 and not re.search(r"\d{2,4}[.,\-/]\d{1,2}", raw):
            out.append(format_phone(raw))
    return dedupe(out)


def strip_phones_from_text(text: str) -> str:
    """Remove embedded phone numbers from an address-like fragment."""
    out = str(text or "")
    for raw_match in [m.group(0) for m in PHONE_BLOCK_RE.finditer(out)]:
        raw = clean_text(raw_match)
        digits = re.sub(r"\D", "", raw)
        if len(digits) in (6, 10, 11):
            out = out.replace(raw_match, " ")
    for raw_match in [m.group(0) for m in PHONE_DIGIT_TOKEN_RE.finditer(out)]:
        digits = re.sub(r"\D", "", raw_match)
        if len(digits) in (6, 10, 11):
            out = out.replace(raw_match, " ")
    for raw_match in [m.group(0) for m in PHONE_FLEX_RE.finditer(out)]:
        digits = re.sub(r"\D", "", raw_match)
        if len(digits) in (6, 10, 11):
            out = out.replace(raw_match, " ")
    return clean_text(out)


def extract_contact(text: str) -> str:
    src = clean_text(text)
    if not src:
        return ""

    names: list[str] = []
    names.extend(match.group(0) for match in INITIALS_RE.finditer(src))
    names.extend(match.group(0) for match in PERSON_RE.finditer(src))
    for match in SINGLE_NAME_RE.finditer(src):
        token = match.group(0)
        if is_russian_name_token(token):
            names.append(token)

    filtered: list[str] = []
    for name in dedupe(names):
        parts = [part for part in name.split() if part]
        parts_low = {part.lower() for part in parts}
        if has_legal_marker(name) or looks_like_address(name):
            continue
        if parts_low & ROLE_WORDS:
            continue
        if any(part.lower() in BUSINESS_HINT_WORDS for part in parts):
            continue
        if not any(is_russian_name_token(part) for part in parts):
            continue
        filtered.append(name)

    if not filtered:
        return ""
    filtered.sort(key=lambda item: (-len(item.split()), len(item)))
    return clean_text(filtered[0])


def extract_address(text: str) -> str:
    src = clean_text(text)
    if not src:
        return ""

    parts = [clean_text(part) for part in re.split(r"[;|]", src) if clean_text(part)]
    for part in parts:
        if not looks_like_address(part):
            continue
        city_match = re.search(r"\bг\.?\s*[А-ЯЁа-яё-]+", part)
        marker = ADDR_TOKEN_RE.search(part)
        street = STREET_WITH_NUMBER_RE.search(part) if has_street_number_candidate(part) else None
        compact = COMPACT_ADDRESS_RE.search(part)
        starts = [m.start() for m in (city_match, marker, street, compact) if m is not None]
        if starts:
            return strip_phones_from_text(part[min(starts):])
        return strip_phones_from_text(part)

    marker = ADDR_TOKEN_RE.search(src)
    if marker:
        tail = clean_text(src[marker.start() :].split(" ; ", 1)[0])
        if re.search(r"\d", tail):
            return strip_phones_from_text(tail)

    street = STREET_WITH_NUMBER_RE.search(src) if has_street_number_candidate(src) else None
    compact = COMPACT_ADDRESS_RE.search(src)
    best = None
    if compact is not None:
        best = compact
    elif street is not None:
        best = street
    if best:
        start = best.start()
        city_matches = list(CITY_RE.finditer(src[:start]))
        city = city_matches[-1].group(0) if city_matches else ""
        tail = strip_phones_from_text(clean_text(src[start:].split(" ; ", 1)[0]))
        if city and city.lower() not in GENERIC_NON_COMPANY_WORDS:
            return clean_text(f"{city}, {tail}")
        return tail
    return ""


def remove_address_from_company(candidate: str) -> str:
    src = clean_text(candidate)
    if not src:
        return ""
    city_match = re.search(r"\bг\.?\s*[А-ЯЁа-яё-]+", src)
    marker = ADDR_TOKEN_RE.search(src)
    street = STREET_WITH_NUMBER_RE.search(src) if has_street_number_candidate(src) else None
    starts = [m.start() for m in (city_match, marker, street) if m is not None and m.start() > 0]
    if starts:
        src = src[: min(starts)]
    chunks = [clean_text(chunk) for chunk in src.split(",") if clean_text(chunk)]
    if len(chunks) > 1 and looks_like_person(chunks[-1]):
        src = ", ".join(chunks[:-1])
    return normalize_company_text(src.strip(" ,;-"))


def split_company_and_address(raw_name: str) -> tuple[str, str]:
    src = clean_text(raw_name)
    if not src:
        return "", ""
    address = extract_address(src)
    company = remove_address_from_company(src)
    return company, address


def looks_like_company(candidate: str, client_kind: str) -> bool:
    text = normalize_company_text(candidate)
    if not text:
        return False
    low = text.lower()
    if any(word in low for word in GENERIC_NON_COMPANY_WORDS):
        return False
    if EMAIL_RE.search(text) or PHONE_BLOCK_RE.search(text):
        return False
    if looks_like_person(text) or looks_like_address(text) or extract_contact(text):
        return False
    if "физ" in client_kind:
        return False
    if has_legal_marker(text):
        return True
    if re.search(r"[A-Za-z]", text) and len(text) >= 3:
        return True
    if any(word in low.split() for word in BUSINESS_HINT_WORDS):
        return True
    return len(text.split()) >= 2 and not any(is_russian_name_token(part) for part in text.split())


def pick_company(full_text: str, raw_name: str, raw_contact: str, client_kind: str) -> str:
    if "физ" in client_kind:
        return ""
    candidates: list[str] = []
    if has_legal_marker(raw_name):
        candidates.append(remove_address_from_company(raw_name))
    name_company, _ = split_company_and_address(raw_name)
    if name_company:
        candidates.append(name_company)
    chunks = [clean_text(item) for item in re.split(r"[;|]", full_text) if clean_text(item)]
    candidates.extend(remove_address_from_company(chunk) for chunk in chunks)
    for candidate in dedupe(candidates):
        if EMAIL_RE.search(candidate) or PHONE_BLOCK_RE.search(candidate):
            continue
        if looks_like_company(candidate, client_kind):
            return normalize_company_text(candidate)
    return ""


def normalize_address(address: str, company: str) -> str:
    addr = clean_text(address)
    comp = normalize_company_text(company)
    if not addr:
        return ""
    if comp:
        addr = re.sub(rf"(?i)^\s*{re.escape(comp)}\s*,?\s*", "", addr)
    return clean_text(addr)


def infer_contact_for_physical_person(raw_name: str, raw_contact: str, full_text: str) -> str:
    for source in (raw_name, raw_contact, full_text):
        name = extract_contact(source)
        if name:
            return name
    return ""


def normalize_row(*, client_kind: str, raw_name: str, raw_contact: str, row_text: str) -> RowNormalizationResult:
    emails = dedupe(extract_emails(row_text) + extract_emails(raw_name) + extract_emails(raw_contact))
    phones = dedupe(extract_phones(raw_contact) + extract_phones(raw_name) + extract_phones(row_text))
    company = pick_company(row_text, raw_name, raw_contact, client_kind)
    _, address_from_name = split_company_and_address(raw_name)
    address = normalize_address(extract_address(row_text) or address_from_name, company)
    contact = extract_contact(raw_contact) or extract_contact(row_text)

    if "физ" in client_kind:
        contact = infer_contact_for_physical_person(raw_name, raw_contact, row_text)
        company = ""
        if not address:
            address = normalize_address(extract_address(raw_name), company)

    if company and (looks_like_person(company) or looks_like_address(company)):
        company = ""

    return RowNormalizationResult(
        company=normalize_company_text(company),
        contact=clean_text(contact),
        address=clean_text(address),
        phones=", ".join(phones),
        emails=", ".join(emails),
    )


def detect_header_row(ws: Worksheet) -> int | None:
    for row_idx in range(1, min(ws.max_row, 40) + 1):
        values = [clean_text(ws.cell(row_idx, col_idx).value).lower() for col_idx in range(1, min(ws.max_column, 16) + 1)]
        row_text = " | ".join(value for value in values if value)
        if "назв" in row_text and ("контакт" in row_text or "телефон" in row_text or "почта" in row_text):
            return row_idx
    return None


def ensure_output_columns(ws: Worksheet, header_row: int) -> dict[str, int]:
    header_to_col: dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        header = clean_text(ws.cell(header_row, col_idx).value)
        if header:
            header_to_col[header] = col_idx
    next_col = ws.max_column + 1
    for header in NEW_COLUMNS:
        if header not in header_to_col:
            ws.cell(header_row, next_col, header)
            header_to_col[header] = next_col
            next_col += 1
    return header_to_col


def infer_source_columns(headers: list[str]) -> tuple[int | None, int | None, int | None]:
    kind_idx: int | None = None
    name_idx: int | None = None
    contact_idx: int | None = None
    for idx, header in enumerate(headers):
        low = header.lower()
        if kind_idx is None and "вид" in low and "клиент" in low:
            kind_idx = idx
        if name_idx is None and "назв" in low:
            name_idx = idx
        if contact_idx is None and ("контакт" in low or "почта" in low or "телефон" in low):
            contact_idx = idx
    return kind_idx, name_idx, contact_idx


def build_dataframe(ws: Worksheet, header_row: int) -> tuple[pd.DataFrame, list[str]]:
    header_values = [clean_text(cell.value) for cell in ws[header_row]]
    headers = [value or f"__col_{idx}" for idx, value in enumerate(header_values, start=1)]
    rows = [list(row) for row in ws.iter_rows(min_row=header_row + 1, values_only=True)]
    return pd.DataFrame(rows, columns=headers), headers


def row_text_from_series(row: pd.Series, headers: list[str]) -> str:
    pieces: list[str] = []
    for header in headers:
        if header in NEW_COLUMNS:
            continue
        value = clean_text(row.get(header, ""))
        if value:
            pieces.append(value)
    return clean_text(" ; ".join(pieces))


def process_sheet(ws: Worksheet) -> int:
    header_row = detect_header_row(ws)
    if header_row is None:
        return 0
    frame, headers = build_dataframe(ws, header_row)
    kind_idx, name_idx, contact_idx = infer_source_columns(headers)
    if name_idx is None:
        return 0

    out_cols = ensure_output_columns(ws, header_row)
    updated = 0
    for frame_row_idx, (_, row) in enumerate(frame.iterrows(), start=header_row + 1):
        row_text = row_text_from_series(row, headers)
        if not row_text:
            continue
        client_kind = clean_text(row.iloc[kind_idx]) if kind_idx is not None else ""
        raw_name = clean_text(row.iloc[name_idx])
        raw_contact = clean_text(row.iloc[contact_idx]) if contact_idx is not None else ""
        result = normalize_row(
            client_kind=client_kind.lower(),
            raw_name=raw_name,
            raw_contact=raw_contact,
            row_text=row_text,
        )
        ws.cell(frame_row_idx, out_cols["Компания (auto)"], result.company)
        ws.cell(frame_row_idx, out_cols["Контактное лицо (auto)"], result.contact)
        ws.cell(frame_row_idx, out_cols["Адрес (auto)"], result.address)
        ws.cell(frame_row_idx, out_cols["Телефоны (auto)"], result.phones)
        ws.cell(frame_row_idx, out_cols["Email (auto)"], result.emails)
        updated += 1
    return updated


def make_backup(src: Path) -> Path:
    backup = src.with_name(src.stem + "_backup_before_normalize" + src.suffix)
    shutil.copy2(src, backup)
    return backup


def main() -> int:
    if len(sys.argv) < 2:
        print("Usage: normalize_client_book.py <input.xlsx> [output.xlsx]")
        return 2
    src = Path(sys.argv[1])
    if not src.exists():
        print(f"Input file not found: {src}")
        return 2
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_name(src.stem + "_structured.xlsx")
    if src.resolve() == dst.resolve():
        backup = make_backup(src)
        print(f"Backup created: {backup}")
    wb = load_workbook(src)
    total = 0
    touched: list[str] = []
    for ws in wb.worksheets:
        processed = process_sheet(ws)
        total += processed
        if processed:
            touched.append(ws.title)
    wb.save(dst)
    print(f"OK: {dst}")
    print(f"Rows processed: {total}")
    print(f"Sheets updated: {', '.join(touched) if touched else '-'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
